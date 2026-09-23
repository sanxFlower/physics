from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path

rows = [
    (1, "组长", "舒雅晴"), (1, "2", "陈家睿"), (1, "3", "付子璇"),
    (1, "4", "何梓涵"), (1, "5", "刘珈汝"), (1, "6", "史子鑫"),
    (1, "7", "吴可"), (1, "8", "徐诗雨"), (1, "9", "姚芊辰"),
    (1, "10", "周禄彦"),
    (2, "组长", "柯锦宏"), (2, "2", "江奕婷"), (2, "3", "焦梓翰"),
    (2, "4", "黎梓颖"), (2, "5", "刘子逸"), (2, "6", "年奕宏"),
    (2, "7", "潘皓然"), (2, "8", "吴景轩"), (2, "9", "张紫妍"),
    (3, "组长", "周元深"), (3, "2", "陈涵阳"), (3, "3", "杜宝涵"),
    (3, "4", "龚烨文"), (3, "5", "贺逸霏"), (3, "6", "雷懿轩"),
    (3, "7", "严秋琴"), (3, "8", "杨雨童"), (3, "9", "张博洋"),
    (4, "组长", "周以彤"), (4, "2", "陈明轩"), (4, "3", "向舒琪"),
    (4, "4", "辛道政"), (4, "5", "杨谨瑞"), (4, "6", "杨煜浠"),
    (4, "7", "岳之松"), (4, "8", "张胡曦"), (4, "9", "张明睿"),
    (5, "组长", "夏未涵"), (5, "2", "鲍皓然"), (5, "3", "管曼舒"),
    (5, "4", "胡一曼"), (5, "5", "黄子轩"), (5, "6", "李国盛"),
    (5, "7", "刘一鸣"), (5, "8", "叶天伦"), (5, "9", "邹玉欣"),
]

out = Path(r"C:\Users\San\workspace\physics\高中物理\光谷二高\26级-高一\学生名单.xlsx")
wb = Workbook()
ws = wb.active
ws.title = "学生名单"
ws.append(["小组", "组内编号", "学生姓名"])
for row in rows:
    ws.append(row)

header_fill = PatternFill("solid", fgColor="1F4E78")
for cell in ws[1]:
    cell.font = Font(color="FFFFFF", bold=True)
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(horizontal="center", vertical="center")

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:C{ws.max_row}"
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 16
ws.row_dimensions[1].height = 22

tab = Table(displayName="学生名单表", ref=f"A1:C{ws.max_row}")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                    showLastColumn=False, showRowStripes=True, showColumnStripes=False)
ws.add_table(tab)
wb.save(out)
print(f"{out}\nrows={len(rows)}")
